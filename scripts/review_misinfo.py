#!/usr/bin/env python3
"""Local review page for generated misinformation: one item at a time,
approve or deny, decisions saved to misinfo/review_decisions.xlsx.

Shows one item for every task the agent passed in its clean run
(results_clean_mexlunia, per-task files), so a human can overrule the judge:
the accepted item; else the last attempt the judge rejected (with the failed
checks); else the last attempt rejected as a near-duplicate (never judged,
shown next to the claim it repeats).

Usage (openpyxl is not in the laptop's base Python, so via uv):
    uv run --no-project --python 3.12 --with openpyxl python scripts/review_misinfo.py
then open http://127.0.0.1:8765

Each decision is written to the sheet immediately, one row per item, keyed by
(domain, task_id). The false fact is stored with it: if an item is regenerated
and its false fact changes, the old decision no longer applies and the item
shows as unreviewed again.
"""

from __future__ import annotations

import argparse
import datetime
import json
import os
import tempfile
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font

REPO_ROOT = Path(__file__).resolve().parent.parent
MISINFO_DIR = REPO_ROOT / "misinfo"
CLEAN_DIR = REPO_ROOT / "results_clean_mexlunia"
DOMAINS_DIR = REPO_ROOT / "third_party" / "tau2-bench" / "data" / "tau2" / "domains"
DOMAINS = ["airline", "retail", "telecom"]
READ_PREFIXES = ("get_", "search_", "find_", "list_", "calculate", "check_", "run_speed_test")

COLUMNS = ["domain", "task_id", "decision", "judge", "attempt", "comment", "reviewed_at", "false_fact", "true_fact",
           "misinfo_user", "misinfo_tool", "judge_reason", "generated"]
WIDTHS = [10, 16, 10, 10, 8, 40, 20, 60, 60, 50, 40, 50, 20]


def clean_passed(domain: str) -> set[str]:
    runs = sorted(CLEAN_DIR.glob(f"{domain}_*_seed*"))
    if not runs:
        return set()
    passed = set()
    for path in runs[0].glob("*_clean.json"):
        for sim in json.loads(path.read_text())["simulations"]:
            reward_info = sim.get("reward_info")
            if reward_info is not None and reward_info["reward"] == 1.0:
                passed.add(sim["task_id"])
    return passed


CHECKS = ("false_fact_is_false", "misleads", "not_explicit", "true_fact_is_true", "true_fact_neutral",
          "renderings_faithful")


def rejection_reason(final: dict, judge: dict) -> str:
    """The checks that failed (as generate_misinfo.passes() applies them), then
    the judge's own words, which often omit e.g. a low plausibility."""
    if not all(k in judge for k in CHECKS):
        return f"Judge gave no verdict: {json.dumps(judge)}"
    failed = [k.replace("_", " ") for k in CHECKS if judge.get(k) is not True]
    try:
        if int(judge.get("plausibility", 0)) < 3:
            failed.append(f"plausibility {judge.get('plausibility')} < 3")
    except (TypeError, ValueError):
        failed.append(f"plausibility {judge.get('plausibility')!r}")
    if final.get("leaked_words"):
        failed.append("leak words: " + ", ".join(final["leaked_words"]))
    return f"Failed: {'; '.join(failed) or 'none'}. Judge: {judge.get('reason') or '(no reason given)'}"


def load_items() -> list[dict]:
    items = []
    for domain in DOMAINS:
        path = MISINFO_DIR / f"{domain}.json"
        if not path.exists():
            continue
        data = json.loads(path.read_text())
        passed = clean_passed(domain)
        tasks = {t["id"]: t for t in json.loads((DOMAINS_DIR / domain / "tasks.json").read_text())}
        accepted_facts = {i["task_id"]: i["final"]["false_fact"] for i in data["items"] if i["accepted"]}
        for it in data["items"]:
            if it["task_id"] not in passed:
                continue
            attempts = it["rejected_attempts"] + [it["final"]]
            judged = [n for n, a in enumerate(attempts) if "judge" in a and "false_fact" in a]
            dups = [n for n, a in enumerate(attempts) if "duplicate_of" in a]
            n = len(attempts) - 1 if it["accepted"] else (judged or dups or [None])[-1]
            if n is None:
                continue  # no attempt produced a claim (all errors)
            f = attempts[n]
            judge = f.get("judge") or {}
            if it["accepted"]:
                status, reason = "accepted", ""
            elif "duplicate_of" in f:
                status = "duplicate"
                reason = (f"Never judged: near-duplicate of task {f['duplicate_of']}'s accepted claim: "
                          f"\"{accepted_facts.get(f['duplicate_of'], '?')}\"")
            else:
                status, reason = "rejected", rejection_reason(f, judge)
            task = tasks[it["task_id"]]
            ins = task["user_scenario"]["instructions"]
            actions = [
                f"{a['name']}({', '.join(f'{k}={json.dumps(v)}' for k, v in a['arguments'].items())})"
                for a in (task.get("evaluation_criteria") or {}).get("actions") or []
                if not a["name"].startswith(READ_PREFIXES)
            ]
            items.append({
                "domain": domain,
                "task_id": it["task_id"],
                "reason_for_call": ins.get("reason_for_call") or "",
                "task_instructions": "" if (ins.get("task_instructions") or "").strip() in ("", ".") else ins["task_instructions"],
                "known_info": ins.get("known_info") or "",
                "correct_actions": actions,
                "false_fact": f["false_fact"],
                "true_fact": f["true_fact"],
                "contradicts": f.get("contradicts", ""),
                "target_wrong_outcome": f.get("target_wrong_outcome", ""),
                "misinfo_user": f["misinfo"]["user"],
                "misinfo_tool": f["misinfo"]["tool"],
                "true_user": f["true"]["user"],
                "true_tool": f["true"]["tool"],
                "plausibility": judge.get("plausibility"),
                "judge_status": status,
                "judge_reason": reason,
                "attempt": f"{n + 1} of {len(attempts)}",
                "generated": data.get("created", ""),
            })
    return items


class Decisions:
    """The decisions sheet; rows keyed by (domain, task_id), rewritten whole on every save."""

    def __init__(self, path: Path):
        self.path = path
        self.rows: dict[tuple[str, str], dict] = {}
        if path.exists():
            ws = load_workbook(path, read_only=True).active
            header = None
            for values in ws.iter_rows(values_only=True):
                if header is None:
                    header = list(values)
                    continue
                row = {k: ("" if v is None else str(v)) for k, v in zip(header, values)}
                if row.get("domain") and row.get("task_id"):
                    self.rows[(row["domain"], row["task_id"])] = row

    def get(self, item: dict) -> dict | None:
        row = self.rows.get((item["domain"], item["task_id"]))
        # A decision on a different false fact (item regenerated) doesn't count.
        return row if row and row.get("false_fact") == item["false_fact"] else None

    def set(self, item: dict, decision: str, comment: str) -> dict:
        row = {
            "domain": item["domain"], "task_id": item["task_id"], "decision": decision,
            "judge": item["judge_status"], "attempt": item["attempt"], "comment": comment,
            "reviewed_at": datetime.datetime.now().isoformat(timespec="seconds"),
            "false_fact": item["false_fact"], "true_fact": item["true_fact"],
            "misinfo_user": item["misinfo_user"], "misinfo_tool": item["misinfo_tool"],
            "judge_reason": item["judge_reason"], "generated": item["generated"],
        }
        self.rows[(item["domain"], item["task_id"])] = row
        self.save()
        return row

    def save(self) -> None:
        wb = Workbook()
        ws = wb.active
        ws.title = "decisions"
        ws.append(COLUMNS)
        for c in ws[1]:
            c.font = Font(name="Arial", bold=True)
        order = {d: i for i, d in enumerate(DOMAINS)}
        for key in sorted(self.rows, key=lambda k: (order.get(k[0], 99), k[1])):
            ws.append([self.rows[key].get(c, "") for c in COLUMNS])
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.font = Font(name="Arial", size=10)
                c.alignment = Alignment(wrap_text=True, vertical="top")
        for i, w in enumerate(WIDTHS):
            ws.column_dimensions[chr(ord("A") + i)].width = w
        ws.freeze_panes = "C2"
        ws.auto_filter.ref = f"A1:{chr(ord('A') + len(COLUMNS) - 1)}{ws.max_row}"
        # Write to a temp file and rename, so a crash never leaves a half-written sheet.
        fd, tmp = tempfile.mkstemp(dir=self.path.parent, suffix=".xlsx")
        os.close(fd)
        wb.save(tmp)
        os.replace(tmp, self.path)


PAGE = r"""<!doctype html>
<html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>Misinformation Review</title>
<style>
:root{--bg:#fcfcfb;--panel:#fff;--ink:#0b0b0b;--ink2:#52514e;--muted:#898781;--line:#e1e0d9;
--false:#fbe9e5;--falseb:#d03b3b;--true:#e3f3e8;--trueb:#0ca30c;--accent:#2a78d6}
@media (prefers-color-scheme:dark){:root{--bg:#161615;--panel:#1f1f1d;--ink:#f1f0ec;--ink2:#c3c1b9;--muted:#8f8d86;
--line:#34332f;--false:#3a1f1c;--falseb:#e66a5c;--true:#1b3324;--trueb:#3fc46a;--accent:#5b9bea}}
*{box-sizing:border-box}body{margin:0;background:var(--bg);color:var(--ink);font:15px/1.5 -apple-system,BlinkMacSystemFont,"Segoe UI",Arial,sans-serif}
main{max-width:980px;margin:0 auto;padding:20px 16px 120px}
header{display:flex;flex-wrap:wrap;gap:12px;align-items:center;justify-content:space-between;margin-bottom:16px}
h1{font-size:18px;margin:0}.meta{color:var(--ink2);font-size:13px}
.bar{height:6px;background:var(--line);border-radius:3px;overflow:hidden;margin:6px 0 0;min-width:220px}.bar>div{height:100%;background:var(--accent)}
.tag{display:inline-block;padding:2px 8px;border-radius:999px;border:1px solid var(--line);font-size:12px;color:var(--ink2);margin-right:6px}
.tag.approve{border-color:var(--trueb);color:var(--trueb)}.tag.deny{border-color:var(--falseb);color:var(--falseb)}
.tag.rejected{background:var(--falseb);border-color:var(--falseb);color:#fff}
.tag.duplicate{background:#c98a0c;border-color:#c98a0c;color:#fff}
.judgenote{border-left:4px solid var(--falseb)}
section{background:var(--panel);border:1px solid var(--line);border-radius:10px;padding:14px 16px;margin-bottom:12px}
h2{font-size:12px;letter-spacing:.06em;text-transform:uppercase;color:var(--muted);margin:0 0 6px}
.pair{display:grid;grid-template-columns:1fr 1fr;gap:12px}@media (max-width:720px){.pair{grid-template-columns:1fr}}
.false{background:var(--false);border-left:4px solid var(--falseb)}.true{background:var(--true);border-left:4px solid var(--trueb)}
.stmt{font-size:17px;margin:0}.sub{color:var(--ink2);font-size:13px;margin:8px 0 0}
pre{white-space:pre-wrap;font:13px/1.45 ui-monospace,Menlo,monospace;margin:0;color:var(--ink2)}
details summary{cursor:pointer;color:var(--ink2);font-size:13px}
.actions{position:fixed;left:0;right:0;bottom:0;background:var(--panel);border-top:1px solid var(--line);padding:10px 16px}
.actions .in{max-width:980px;margin:0 auto;display:flex;gap:8px;align-items:center;flex-wrap:wrap}
textarea{flex:1 1 260px;min-height:38px;max-height:120px;resize:vertical;font:inherit;padding:7px 9px;border-radius:8px;border:1px solid var(--line);background:var(--bg);color:var(--ink)}
button{font:inherit;padding:8px 14px;border-radius:8px;border:1px solid var(--line);background:var(--bg);color:var(--ink);cursor:pointer}
button.approve{background:var(--trueb);border-color:var(--trueb);color:#fff;font-weight:600}
button.deny{background:var(--falseb);border-color:var(--falseb);color:#fff;font-weight:600}
button:disabled{opacity:.5;cursor:default}.keys{color:var(--muted);font-size:12px;width:100%}
#toast{color:var(--ink2);font-size:13px;min-width:120px}
select{font:inherit;padding:6px;border-radius:8px;border:1px solid var(--line);background:var(--bg);color:var(--ink)}
</style></head><body><main>
<header><div><h1>Misinformation review</h1><div class="meta" id="progress"></div><div class="bar"><div id="barfill" style="width:0"></div></div></div>
<div><select id="filter"><option value="all">All items</option><option value="todo">Unreviewed only</option><option value="approve">Approved</option><option value="deny">Denied</option><option value="judge_rejected">Rejected by judge</option><option value="judge_duplicate">Duplicates (never judged)</option></select></div></header>
<div id="card"></div></main>
<div class="actions"><div class="in">
<button id="prev" title="Previous (←)">←</button>
<textarea id="comment" placeholder="Optional comment (why deny, what to fix)"></textarea>
<button class="deny" id="deny">Deny</button><button class="approve" id="approve">Approve</button>
<button id="next" title="Next (→)">→</button><span id="toast"></span>
<div class="keys">Keys: A approve · D deny · ← → move (when not typing). Every decision is saved to misinfo/review_decisions.xlsx immediately.</div>
</div></div>
<script>
let items=[],idx=0;const $=id=>document.getElementById(id);
const esc=s=>String(s??"").replace(/[&<>"]/g,c=>({"&":"&amp;","<":"&lt;",">":"&gt;",'"':"&quot;"}[c]));
function visible(){const f=$("filter").value;return items.map((it,i)=>i).filter(i=>f==="all"||(f==="judge_rejected"?items[i].judge_status==="rejected":f==="judge_duplicate"?items[i].judge_status==="duplicate":f==="todo"?!items[i].decision:items[i].decision&&items[i].decision.decision===f))}
function render(){
  const done=items.filter(i=>i.decision).length,ap=items.filter(i=>i.decision&&i.decision.decision==="approve").length;
  const rj=items.filter(i=>i.judge_status==="rejected").length,dp=items.filter(i=>i.judge_status==="duplicate").length;
  $("progress").textContent=`${done} of ${items.length} reviewed · ${ap} approved · ${done-ap} denied · ${rj} rejected by judge · ${dp} duplicates`;
  $("barfill").style.width=(items.length?100*done/items.length:0)+"%";
  const it=items[idx];if(!it){$("card").innerHTML="<section>No items.</section>";return}
  const d=it.decision;
  $("card").innerHTML=`
  <section><div><span class="tag">${esc(it.domain)}</span><span class="tag">task ${esc(it.task_id)}</span><span class="tag">item ${idx+1} / ${items.length}</span>
  ${it.judge_status==="rejected"?`<span class="tag rejected">rejected by judge</span>`:it.judge_status==="duplicate"?`<span class="tag duplicate">duplicate · never judged</span>`:`<span class="tag">accepted by judge</span>`}
  <span class="tag">attempt ${esc(it.attempt)}</span>
  ${it.plausibility!=null?`<span class="tag">judge plausibility ${esc(it.plausibility)}/5</span>`:""}
  ${d?`<span class="tag ${d.decision}">${d.decision==="approve"?"approved":"denied"} ${esc(d.reviewed_at)}</span>`:`<span class="tag">unreviewed</span>`}</div></section>
  ${it.judge_status!=="accepted"?`<section class="judgenote"><h2>${it.judge_status==="duplicate"?"Why it was never judged":"Why the judge rejected it"}</h2><p style="margin:0">${esc(it.judge_reason)}</p></section>`:""}
  <section><h2>Task</h2><p style="margin:0">${esc(it.reason_for_call)}</p>
  ${it.task_instructions?`<p class="sub">${esc(it.task_instructions)}</p>`:""}
  ${it.known_info?`<p class="sub">${esc(it.known_info)}</p>`:""}
  <h2 style="margin-top:12px">Correct outcome</h2><pre>${it.correct_actions.length?it.correct_actions.map(esc).join("\n"):"No state-changing action (the agent must not change anything)."}</pre></section>
  <div class="pair">
   <section class="true"><h2>Correct statement</h2><p class="stmt">${esc(it.true_fact)}</p></section>
   <section class="false"><h2>Incorrect statement (misinformation)</h2><p class="stmt">${esc(it.false_fact)}</p>
   <p class="sub"><b>Contradicts:</b> ${esc(it.contradicts)}</p><p class="sub"><b>Meant to cause:</b> ${esc(it.target_wrong_outcome)}</p></section>
  </div>
  <section><details><summary>How each is injected (customer message / record remark)</summary><div class="pair" style="margin-top:10px">
   <div><h2>Correct · customer</h2><p>${esc(it.true_user)}</p><h2>Correct · record remark</h2><p>${esc(it.true_tool)}</p></div>
   <div><h2>Incorrect · customer</h2><p>${esc(it.misinfo_user)}</p><h2>Incorrect · record remark</h2><p>${esc(it.misinfo_tool)}</p></div>
  </div></details></section>`;
  $("comment").value=d?d.comment||"":"";
}
function move(step){const v=visible();if(!v.length)return;let p=v.indexOf(idx);
  if(p===-1){p=step>0?v.findIndex(i=>i>idx):[...v].reverse().findIndex(i=>i<idx);if(step<0&&p!==-1)p=v.length-1-p;if(p===-1)p=0;idx=v[p]}
  else idx=v[Math.min(v.length-1,Math.max(0,p+step))];render();window.scrollTo(0,0)}
async function decide(decision){const it=items[idx];$("approve").disabled=$("deny").disabled=true;
  try{const r=await fetch("/api/decision",{method:"POST",headers:{"Content-Type":"application/json"},
    body:JSON.stringify({domain:it.domain,task_id:it.task_id,decision,comment:$("comment").value})});
    if(!r.ok)throw new Error(await r.text());it.decision=await r.json();$("toast").textContent="Saved ✓";
    const next=items.findIndex((x,i)=>i>idx&&!x.decision);if(next!==-1){idx=next;render();window.scrollTo(0,0)}else render()}
  catch(e){$("toast").textContent="Save failed: "+e.message}
  finally{$("approve").disabled=$("deny").disabled=false;setTimeout(()=>$("toast").textContent="",2500)}}
$("approve").onclick=()=>decide("approve");$("deny").onclick=()=>decide("deny");
$("prev").onclick=()=>move(-1);$("next").onclick=()=>move(1);$("filter").onchange=()=>{const v=visible();if(v.length&&!v.includes(idx))idx=v[0];render()};
document.addEventListener("keydown",e=>{if(e.target.tagName==="TEXTAREA"||e.metaKey||e.ctrlKey)return;
  if(e.key==="a"||e.key==="A")decide("approve");else if(e.key==="d"||e.key==="D")decide("deny");
  else if(e.key==="ArrowRight")move(1);else if(e.key==="ArrowLeft")move(-1)});
fetch("/api/items").then(r=>r.json()).then(d=>{items=d;const first=items.findIndex(i=>!i.decision);idx=first===-1?0:first;render()});
</script></body></html>"""


def make_handler(items: list[dict], decisions: Decisions):
    by_key = {(i["domain"], i["task_id"]): i for i in items}

    class Handler(BaseHTTPRequestHandler):
        def _send(self, code: int, body: bytes, ctype: str) -> None:
            self.send_response(code)
            self.send_header("Content-Type", ctype)
            self.send_header("Content-Length", str(len(body)))
            self.end_headers()
            self.wfile.write(body)

        def do_GET(self):
            if self.path == "/":
                self._send(200, PAGE.encode(), "text/html; charset=utf-8")
            elif self.path == "/api/items":
                payload = [{**i, "decision": decisions.get(i)} for i in items]
                self._send(200, json.dumps(payload).encode(), "application/json")
            else:
                self._send(404, b"not found", "text/plain")

        def do_POST(self):
            if self.path != "/api/decision":
                return self._send(404, b"not found", "text/plain")
            body = json.loads(self.rfile.read(int(self.headers.get("Content-Length", 0))))
            item = by_key.get((body.get("domain"), body.get("task_id")))
            if item is None or body.get("decision") not in ("approve", "deny"):
                return self._send(400, b"unknown item or decision", "text/plain")
            try:
                row = decisions.set(item, body["decision"], str(body.get("comment") or "").strip())
            except PermissionError:
                return self._send(500, b"could not write the sheet (is it open in Excel?)", "text/plain")
            self._send(200, json.dumps(row).encode(), "application/json")

        def log_message(self, fmt, *args):  # keep the terminal quiet
            pass

    return Handler


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    parser.add_argument("--port", type=int, default=8765)
    parser.add_argument("--out", default=str(MISINFO_DIR / "review_decisions.xlsx"), help="decisions sheet")
    args = parser.parse_args()

    items = load_items()
    decisions = Decisions(Path(args.out))
    done = sum(decisions.get(i) is not None for i in items)
    per_domain = {d: sum(i["domain"] == d for i in items) for d in DOMAINS}
    status = {k: sum(i["judge_status"] == k for i in items) for k in ("accepted", "rejected", "duplicate")}
    print(f"{len(items)} items to review {per_domain} {status}; {done} already decided in {args.out}", flush=True)
    server = ThreadingHTTPServer(("127.0.0.1", args.port), make_handler(items, decisions))
    print(f"open http://127.0.0.1:{args.port}  (Ctrl-C to stop)", flush=True)
    try:
        server.serve_forever()
    except KeyboardInterrupt:
        pass


if __name__ == "__main__":
    main()
