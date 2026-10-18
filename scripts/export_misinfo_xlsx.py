#!/usr/bin/env python3
"""Write misinfo/misinfo_side_by_side.xlsx: every generated item with its false
and true versions in adjacent columns, one sheet per domain, for reading and
human review.

Usage (openpyxl is not in the laptop's base Python, so via uv):
    uv run --no-project --with openpyxl python scripts/export_misinfo_xlsx.py

Reads misinfo/<domain>.json (scripts/generate_misinfo.py) and the clean runs'
per-task files in results_clean_mexlunia/ (for the "clean passed" column).
"""

from __future__ import annotations

import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

REPO_ROOT = Path(__file__).resolve().parent.parent
MISINFO_DIR = REPO_ROOT / "misinfo"
CLEAN_DIR = REPO_ROOT / "results_clean_mexlunia"
OUT_PATH = MISINFO_DIR / "misinfo_side_by_side.xlsx"
DOMAINS = ["airline", "retail", "telecom"]

FONT = "Arial"
FALSE_FILL = PatternFill("solid", fgColor="F8DDD8")  # light red: false (misinformation) columns
TRUE_FILL = PatternFill("solid", fgColor="DCEFE3")  # light green: true (control) columns
INFO_FILL = PatternFill("solid", fgColor="EEEDE8")  # neutral: task / judge columns
THIN = Side(style="thin", color="C8C6BE")

# (header, width, group); group picks the header fill. False/true pairs sit side by side.
COLUMNS = [
    ("Task", 14, "info"),
    ("Target record(s)", 14, "info"),
    ("Accepted by judge", 10, "info"),
    ("Clean run passed", 10, "info"),
    ("Attempts", 9, "info"),
    ("Plausibility (1-5)", 11, "info"),
    ("Contradicts (policy / record)", 40, "info"),
    ("Target wrong outcome", 36, "info"),
    ("FALSE fact", 48, "false"),
    ("TRUE fact", 48, "true"),
    ("FALSE - said by customer", 48, "false"),
    ("TRUE - said by customer", 48, "true"),
    ("FALSE - record remark", 40, "false"),
    ("TRUE - record remark", 40, "true"),
    ("Judge reason / error", 48, "info"),
]
FILLS = {"info": INFO_FILL, "false": FALSE_FILL, "true": TRUE_FILL}


def clean_passed(domain: str) -> dict[str, bool]:
    runs = sorted(CLEAN_DIR.glob(f"{domain}_*_seed*"))
    if not runs:
        return {}
    out = {}
    for path in sorted(runs[0].glob("*_clean.json")):
        for sim in json.loads(path.read_text())["simulations"]:
            reward_info = sim.get("reward_info")
            out[sim["task_id"]] = reward_info is not None and reward_info["reward"] == 1.0
    return out


def row_for(it: dict, passed: dict[str, bool]) -> list:
    f = it["final"]
    judge = f.get("judge") or {}
    if "duplicate_of" in f:
        reason = f"near-duplicate of task {f['duplicate_of']}"
    else:
        reason = judge.get("reason") if "reason" in judge else f.get("error") or (json.dumps(judge) if judge else "")
    tid = it["task_id"]
    return [
        tid,
        ", ".join(it["target_ids"]),
        "yes" if it["accepted"] else "no",
        {True: "yes", False: "no"}.get(passed.get(tid), "n/a"),
        len(it["rejected_attempts"]) + 1,
        judge.get("plausibility"),
        f.get("contradicts"),
        f.get("target_wrong_outcome"),
        f.get("false_fact"),
        f.get("true_fact"),
        (f.get("misinfo") or {}).get("user"),
        (f.get("true") or {}).get("user"),
        (f.get("misinfo") or {}).get("tool"),
        (f.get("true") or {}).get("tool"),
        reason,
    ]


def write_sheet(ws, rows: list[list]) -> None:
    for c, (header, width, group) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=c, value=header)
        cell.font = Font(name=FONT, bold=True)
        cell.fill = FILLS[group]
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = Border(bottom=THIN)
        ws.column_dimensions[get_column_letter(c)].width = width
    for r, values in enumerate(rows, start=2):
        for c, v in enumerate(values, start=1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.font = Font(name=FONT, size=10)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.border = Border(bottom=THIN)
            group = COLUMNS[c - 1][2]
            if group != "info":
                cell.fill = FILLS[group]
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = "B2"  # header row and task column stay visible
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{len(rows) + 1}"


def main() -> None:
    wb = Workbook()
    readme = wb.active
    readme.title = "README"
    notes = [
        "Misinformation items, false and true versions side by side",
        "",
        "One sheet per domain, one row per task (its final generation attempt).",
        "Red columns: the false fact (misinformation) and its two renderings. Green columns: the matched true fact (control).",
        "'said by customer' = user-channel injection; 'record remark' = tool-channel injection (remarks field on the looked-up record).",
        "Accepted by judge: gpt-oss:120b passed all checks with plausibility >= 3 and no leak words.",
        "Clean run passed: the agent passed this task without misinformation (results_clean_mexlunia, seed 42). "
        "Only accepted items on clean-passed tasks can measure a misinformation effect.",
        "Use the column filters to show e.g. only accepted = yes and clean run passed = yes.",
        "",
        "Generated by scripts/export_misinfo_xlsx.py from misinfo/<domain>.json; rerun it after regenerating misinformation.",
    ]
    for r, text in enumerate(notes, start=1):
        cell = readme.cell(row=r, column=1, value=text)
        cell.font = Font(name=FONT, bold=(r == 1), size=12 if r == 1 else 10)
    readme.column_dimensions["A"].width = 140

    for domain in DOMAINS:
        path = MISINFO_DIR / f"{domain}.json"
        if not path.exists():
            continue
        data = json.loads(path.read_text())
        passed = clean_passed(domain)
        rows = [row_for(it, passed) for it in data["items"]]
        write_sheet(wb.create_sheet(domain), rows)
        print(f"{domain}: {len(rows)} items ({sum(r[2] == 'yes' for r in rows)} accepted)")

    wb.save(OUT_PATH)
    print(f"wrote {OUT_PATH.relative_to(REPO_ROOT)}")


if __name__ == "__main__":
    main()
